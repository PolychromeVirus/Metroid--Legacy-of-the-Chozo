#!/usr/bin/env python3
"""Export Legacy of the Chozo card data from Excel to GameMaker-friendly JSON."""

from __future__ import annotations

import argparse
import json
import re
import sys
from collections import Counter
from dataclasses import asdict, dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

try:
    from openpyxl import load_workbook
except ImportError:
    print(
        "Missing dependency: openpyxl. Install it with "
        "`python -m pip install -r tools/requirements.txt`.",
        file=sys.stderr,
    )
    raise SystemExit(2)


SCHEMA_VERSION = 1
CARD_SHEETS = {
    "Cards": {"pool": "loc", "set": "LOC", "filename": "cards_loc.json"},
    "Set2": {"pool": "lop", "set": "LOP", "filename": "cards_lop.json"},
    "Starter": {
        "pool": "starter",
        "set": "LOC",
        "filename": "cards_starter.json",
    },
}
METROID_SHEET = "Metroids"
CARD_TYPES = {"Character", "Event", "Location", "Relic", "Ship"}
FACTION_CODES = ("BH", "CZ", "GF", "NA", "SP", "PZ")
KNOWN_MARKUP = {"cost", "ex", "des"}
REQUIRED_CARD_COLUMNS = (
    "count",
    "set",
    "faction",
    "type",
    "image",
    "name",
    "tags",
    "effect",
    "strength",
    "deploy",
    "reserve",
    "flavor",
    "allowed_layout",
)
REQUIRED_METROID_COLUMNS = (
    "count",
    "set",
    "alias",
    "name",
    "stage",
    "hazard",
    "vp",
    "effect",
    "flavour",
    "phazon",
)


@dataclass(frozen=True)
class Issue:
    severity: str
    code: str
    sheet: str
    row: int | None
    field: str | None
    message: str


class Validator:
    def __init__(self) -> None:
        self.issues: list[Issue] = []

    def add(
        self,
        severity: str,
        code: str,
        sheet: str,
        row: int | None,
        field: str | None,
        message: str,
    ) -> None:
        self.issues.append(Issue(severity, code, sheet, row, field, message))

    def error(
        self, code: str, sheet: str, row: int | None, field: str | None, message: str
    ) -> None:
        self.add("error", code, sheet, row, field, message)

    def warning(
        self, code: str, sheet: str, row: int | None, field: str | None, message: str
    ) -> None:
        self.add("warning", code, sheet, row, field, message)


def normalize_header(value: Any) -> str:
    return str(value).strip().lower() if value is not None else ""


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).replace("\\n", "\n").strip()


def nullable_text(value: Any) -> str | None:
    text = normalize_text(value)
    return text if text else None


def slugify(value: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "_", value.casefold()).strip("_")
    return slug or "unnamed"


def parse_nonnegative_int(
    value: Any,
    validator: Validator,
    sheet: str,
    row: int,
    field: str,
    *,
    required: bool,
) -> int | None:
    if value is None or value == "":
        if required:
            validator.error(
                "missing_number", sheet, row, field, f"{field} is required."
            )
        return None
    if isinstance(value, bool):
        validator.error(
            "invalid_number", sheet, row, field, f"{field} cannot be boolean."
        )
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        validator.error(
            "invalid_number",
            sheet,
            row,
            field,
            f"{field} must be a nonnegative integer, got {value!r}.",
        )
        return None
    if not number.is_integer() or number < 0:
        validator.error(
            "invalid_number",
            sheet,
            row,
            field,
            f"{field} must be a nonnegative integer, got {value!r}.",
        )
        return None
    return int(number)


def parse_numeric_or_x(
    value: Any, validator: Validator, sheet: str, row: int
) -> int | float | str | None:
    if value is None or value == "":
        return None
    if isinstance(value, str):
        text = value.strip()
        if text.casefold() == "x":
            return "X"
        try:
            value = float(text)
        except ValueError:
            validator.error(
                "invalid_stat",
                sheet,
                row,
                "strength",
                f"strength/security must be numeric, X, or blank; got {text!r}.",
            )
            return text
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return int(value) if float(value).is_integer() else float(value)
    validator.error(
        "invalid_stat",
        sheet,
        row,
        "strength",
        f"strength/security has unsupported value {value!r}.",
    )
    return str(value)


def parse_factions(
    value: Any, validator: Validator, sheet: str, row: int
) -> list[str]:
    raw = normalize_text(value).upper()
    if not raw:
        validator.error("missing_faction", sheet, row, "faction", "faction is required.")
        return []
    factions: list[str] = []
    remaining = raw
    while remaining:
        match = next((code for code in FACTION_CODES if remaining.startswith(code)), None)
        if match is None:
            validator.error(
                "unknown_faction",
                sheet,
                row,
                "faction",
                f"Cannot parse faction code {raw!r}; stopped at {remaining!r}.",
            )
            return factions
        factions.append(match)
        remaining = remaining[len(match) :]
    if len(factions) != len(set(factions)):
        validator.warning(
            "duplicate_faction",
            sheet,
            row,
            "faction",
            f"Faction code {raw!r} repeats a faction.",
        )
    return factions


def parse_tags(value: Any) -> list[str]:
    text = normalize_text(value)
    if not text:
        return []
    return [part.strip() for part in re.split(r"\s+-\s+", text) if part.strip()]


def validate_effect_markup(
    effect: str | None, validator: Validator, sheet: str, row: int
) -> None:
    if not effect:
        return
    for match in re.finditer(r"@\[(.*?)\]", effect):
        payload = match.group(1)
        key, _, argument = payload.partition(",")
        if key not in KNOWN_MARKUP:
            validator.warning(
                "unknown_markup",
                sheet,
                row,
                "effect",
                f"Unknown effect markup key {key!r}.",
            )
        elif key == "cost" and (not argument.isdigit() or int(argument) < 0):
            validator.error(
                "invalid_cost_markup",
                sheet,
                row,
                "effect",
                f"Cost markup must be @[cost,n] with a nonnegative integer: {match.group(0)!r}.",
            )
        elif key in {"ex", "des"} and argument:
            validator.warning(
                "unexpected_markup_argument",
                sheet,
                row,
                "effect",
                f"{match.group(0)!r} does not take an argument.",
            )
    if effect.count("@[") != len(re.findall(r"@\[[^\]]*\]", effect)):
        validator.error(
            "malformed_markup",
            sheet,
            row,
            "effect",
            "Effect contains an unclosed @[...] markup token.",
        )


def build_asset_stems(asset_root: Path) -> set[str]:
    extensions = {".png", ".jpg", ".jpeg", ".webp", ".avif", ".svg"}
    stems: set[str] = set()
    for path in asset_root.rglob("*"):
        if not path.is_file() or path.suffix.casefold() not in extensions:
            continue
        stem = path.stem.casefold()
        stems.add(stem)
        if stem.startswith("card_"):
            stems.add(stem.removeprefix("card_"))
    return stems


def build_headers(
    sheet: Any,
    expected: Iterable[str],
    validator: Validator,
    *,
    starter_count_recovery: bool = False,
) -> dict[str, int]:
    values = [normalize_header(cell.value) for cell in sheet[1]]
    if starter_count_recovery and values and values[0] == "":
        values[0] = "count"
        validator.warning(
            "recovered_header",
            sheet.title,
            1,
            "count",
            "Starter!A1 is blank; column A was interpreted as count.",
        )
    headers: dict[str, int] = {}
    for index, name in enumerate(values):
        if not name:
            continue
        if name in headers:
            validator.error(
                "duplicate_header",
                sheet.title,
                1,
                name,
                f"Header {name!r} appears more than once.",
            )
        headers[name] = index
    for name in expected:
        if name not in headers:
            validator.error(
                "missing_header",
                sheet.title,
                1,
                name,
                f"Required column {name!r} is missing.",
            )
    return headers


def cell_value(row: tuple[Any, ...], headers: dict[str, int], name: str) -> Any:
    index = headers.get(name)
    return row[index] if index is not None and index < len(row) else None


def export_card_sheet(
    sheet: Any,
    config: dict[str, str],
    validator: Validator,
    asset_stems: set[str],
) -> list[dict[str, Any]]:
    headers = build_headers(
        sheet,
        REQUIRED_CARD_COLUMNS,
        validator,
        starter_count_recovery=sheet.title == "Starter",
    )
    cards: list[dict[str, Any]] = []
    seen_ids: set[str] = set()
    seen_names: set[str] = set()

    for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
        name = normalize_text(cell_value(row, headers, "name"))
        populated = any(value is not None and value != "" for value in row)
        if not name:
            if populated:
                validator.warning(
                    "ignored_partial_row",
                    sheet.title,
                    row_number,
                    "name",
                    "Row contains data but has no name and was ignored.",
                )
            continue

        card_id = f"{config['pool']}.{slugify(name)}"
        if card_id in seen_ids:
            validator.error(
                "duplicate_id",
                sheet.title,
                row_number,
                "name",
                f"Generated id {card_id!r} is duplicated in this pool.",
            )
        seen_ids.add(card_id)
        folded_name = name.casefold()
        if folded_name in seen_names:
            validator.error(
                "duplicate_name",
                sheet.title,
                row_number,
                "name",
                f"Card name {name!r} is duplicated in this pool.",
            )
        seen_names.add(folded_name)

        count = parse_nonnegative_int(
            cell_value(row, headers, "count"),
            validator,
            sheet.title,
            row_number,
            "count",
            required=True,
        )
        if count == 0:
            validator.warning(
                "zero_count",
                sheet.title,
                row_number,
                "count",
                "Card has a count of zero and will never enter its pool.",
            )

        set_code = normalize_text(cell_value(row, headers, "set")).upper()
        if set_code != config["set"]:
            validator.error(
                "unexpected_set",
                sheet.title,
                row_number,
                "set",
                f"Expected set {config['set']!r}, got {set_code!r}.",
            )

        card_type = normalize_text(cell_value(row, headers, "type")).title()
        if card_type not in CARD_TYPES:
            validator.error(
                "unknown_card_type",
                sheet.title,
                row_number,
                "type",
                f"Unknown card type {card_type!r}.",
            )

        image = normalize_text(cell_value(row, headers, "image"))
        if not image:
            validator.warning(
                "missing_image",
                sheet.title,
                row_number,
                "image",
                "Card has no image key.",
            )
        elif image.casefold() not in asset_stems:
            validator.warning(
                "unmatched_image",
                sheet.title,
                row_number,
                "image",
                f"No card image with stem {image!r} was found under datafiles/cards.",
            )

        effect = nullable_text(cell_value(row, headers, "effect"))
        validate_effect_markup(effect, validator, sheet.title, row_number)
        stat_value = parse_numeric_or_x(
            cell_value(row, headers, "strength"), validator, sheet.title, row_number
        )
        stat_kind = "security" if card_type == "Ship" else (
            "strength" if card_type == "Character" else None
        )

        cards.append(
            {
                "definition_id": card_id,
                "set": set_code,
                "pool": config["pool"],
                "count": count,
                "name": name,
                "type": card_type.casefold(),
                "faction_code": normalize_text(
                    cell_value(row, headers, "faction")
                ).upper(),
                "factions": parse_factions(
                    cell_value(row, headers, "faction"),
                    validator,
                    sheet.title,
                    row_number,
                ),
                "tags": parse_tags(cell_value(row, headers, "tags")),
                "image": image or None,
                "effect": effect,
                "stat": {"kind": stat_kind, "value": stat_value},
                "costs": {
                    "deploy": parse_nonnegative_int(
                        cell_value(row, headers, "deploy"),
                        validator,
                        sheet.title,
                        row_number,
                        "deploy",
                        required=False,
                    ),
                    "reserve": parse_nonnegative_int(
                        cell_value(row, headers, "reserve"),
                        validator,
                        sheet.title,
                        row_number,
                        "reserve",
                        required=False,
                    ),
                },
                "flavor": nullable_text(cell_value(row, headers, "flavor")),
                "dual": normalize_text(
                    cell_value(row, headers, "dual")
                ).casefold()
                == "x",
                "layout": normalize_text(
                    cell_value(row, headers, "allowed_layout")
                ).casefold()
                or None,
                "source": {"sheet": sheet.title, "row": row_number},
            }
        )

    return cards


def export_metroids(
    sheet: Any, validator: Validator
) -> list[dict[str, Any]]:
    headers = build_headers(sheet, REQUIRED_METROID_COLUMNS, validator)
    metroids: list[dict[str, Any]] = []
    seen_ids: set[str] = set()
    for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
        name = normalize_text(cell_value(row, headers, "name"))
        populated = any(value is not None and value != "" for value in row)
        if not name:
            if populated:
                validator.warning(
                    "ignored_partial_row",
                    sheet.title,
                    row_number,
                    "name",
                    "Row contains data but has no name and was ignored.",
                )
            continue
        metroid_id = f"metroid.{slugify(name)}"
        if metroid_id in seen_ids:
            validator.error(
                "duplicate_id",
                sheet.title,
                row_number,
                "name",
                f"Generated id {metroid_id!r} is duplicated.",
            )
        seen_ids.add(metroid_id)

        stage_raw = cell_value(row, headers, "stage")
        if isinstance(stage_raw, (int, float)) and not isinstance(stage_raw, bool):
            stage: int | None = int(stage_raw)
            stage_key = str(stage)
        else:
            stage = None
            stage_key = normalize_text(stage_raw).casefold()
            if not stage_key:
                validator.error(
                    "missing_stage",
                    sheet.title,
                    row_number,
                    "stage",
                    "Metroid stage is required.",
                )

        count = parse_nonnegative_int(
            cell_value(row, headers, "count"),
            validator,
            sheet.title,
            row_number,
            "count",
            required=True,
        )
        hazard = parse_nonnegative_int(
            cell_value(row, headers, "hazard"),
            validator,
            sheet.title,
            row_number,
            "hazard",
            required=True,
        )
        vp_raw = cell_value(row, headers, "vp")
        try:
            research_value = float(vp_raw)
            if research_value.is_integer():
                research_value = int(research_value)
            if research_value < 0:
                raise ValueError
        except (TypeError, ValueError):
            validator.error(
                "invalid_research_value",
                sheet.title,
                row_number,
                "vp",
                f"VP must be a nonnegative number, got {vp_raw!r}.",
            )
            research_value = None

        metroids.append(
            {
                "definition_id": metroid_id,
                "set": normalize_text(cell_value(row, headers, "set")).upper(),
                "count": count,
                "alias": normalize_text(cell_value(row, headers, "alias")),
                "name": name,
                "stage": stage,
                "stage_key": stage_key,
                "hazard": hazard,
                "research_value": research_value,
                "effect": nullable_text(cell_value(row, headers, "effect")),
                "flavor": nullable_text(cell_value(row, headers, "flavour")),
                "phazon": normalize_text(
                    cell_value(row, headers, "phazon")
                ).casefold()
                == "x",
                "source": {"sheet": sheet.title, "row": row_number},
            }
        )
    return metroids


def payload(
    *,
    source: str,
    sheet: str,
    pool: str,
    records_key: str,
    records: list[dict[str, Any]],
) -> dict[str, Any]:
    total_copies = sum(record.get("count") or 0 for record in records)
    return {
        "schema_version": SCHEMA_VERSION,
        "generated_from": source,
        "sheet": sheet,
        "pool": pool,
        "unique_count": len(records),
        "copy_count": total_copies,
        records_key: records,
    }


def write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(
        json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    temporary.replace(path)


def report_payload(
    source: str,
    outputs: list[dict[str, Any]],
    validator: Validator,
) -> dict[str, Any]:
    severity_counts = Counter(issue.severity for issue in validator.issues)
    return {
        "schema_version": SCHEMA_VERSION,
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "source": source,
        "status": "failed" if severity_counts["error"] else "passed",
        "error_count": severity_counts["error"],
        "warning_count": severity_counts["warning"],
        "outputs": outputs,
        "issues": [asdict(issue) for issue in validator.issues],
    }


def write_markdown_report(path: Path, report: dict[str, Any]) -> None:
    lines = [
        "# Card Data Validation",
        "",
        f"- Status: **{report['status'].upper()}**",
        f"- Source: `{report['source']}`",
        f"- Errors: {report['error_count']}",
        f"- Warnings: {report['warning_count']}",
        "",
        "## Exported pools",
        "",
        "| Pool | Worksheet | Unique definitions | Copies | File |",
        "|---|---|---:|---:|---|",
    ]
    for output in report["outputs"]:
        lines.append(
            f"| {output['pool']} | {output['sheet']} | "
            f"{output['unique_count']} | {output['copy_count']} | "
            f"`{output['filename']}` |"
        )
    lines.extend(["", "## Issues", ""])
    if not report["issues"]:
        lines.append("No validation issues.")
    else:
        lines.extend(
            [
                "| Severity | Location | Code | Message |",
                "|---|---|---|---|",
            ]
        )
        for issue in report["issues"]:
            location = issue["sheet"]
            if issue["row"] is not None:
                location += f"!{issue['row']}"
            if issue["field"]:
                location += f" ({issue['field']})"
            message = issue["message"].replace("|", "\\|").replace("\n", " ")
            lines.append(
                f"| {issue['severity']} | {location} | "
                f"`{issue['code']}` | {message} |"
            )
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> int:
    repo_root = Path(__file__).resolve().parents[1]
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--workbook",
        type=Path,
        default=repo_root / "chozoreference.xlsx",
        help="Path to the canonical .xlsx workbook.",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=repo_root / "datafiles" / "generated",
        help="Directory for generated JSON and validation reports.",
    )
    args = parser.parse_args()
    workbook_path = args.workbook.resolve()
    output_dir = args.output.resolve()

    if not workbook_path.is_file():
        print(f"Workbook not found: {workbook_path}", file=sys.stderr)
        return 2

    try:
        source_label = workbook_path.relative_to(repo_root).as_posix()
    except ValueError:
        source_label = str(workbook_path)

    workbook = load_workbook(workbook_path, data_only=False, read_only=True)
    validator = Validator()
    asset_stems = build_asset_stems(repo_root / "datafiles" / "cards")
    exported: list[tuple[str, dict[str, Any]]] = []
    summaries: list[dict[str, Any]] = []

    for sheet_name, config in CARD_SHEETS.items():
        if sheet_name not in workbook.sheetnames:
            validator.error(
                "missing_sheet",
                sheet_name,
                None,
                None,
                f"Required worksheet {sheet_name!r} is missing.",
            )
            continue
        cards = export_card_sheet(
            workbook[sheet_name], config, validator, asset_stems
        )
        data = payload(
            source=source_label,
            sheet=sheet_name,
            pool=config["pool"],
            records_key="cards",
            records=cards,
        )
        exported.append((config["filename"], data))
        summaries.append(
            {
                "pool": config["pool"],
                "sheet": sheet_name,
                "filename": config["filename"],
                "unique_count": data["unique_count"],
                "copy_count": data["copy_count"],
            }
        )

    if METROID_SHEET not in workbook.sheetnames:
        validator.error(
            "missing_sheet",
            METROID_SHEET,
            None,
            None,
            f"Required worksheet {METROID_SHEET!r} is missing.",
        )
    else:
        metroids = export_metroids(workbook[METROID_SHEET], validator)
        data = payload(
            source=source_label,
            sheet=METROID_SHEET,
            pool="metroids",
            records_key="metroids",
            records=metroids,
        )
        filename = "metroids.json"
        exported.append((filename, data))
        summaries.append(
            {
                "pool": "metroids",
                "sheet": METROID_SHEET,
                "filename": filename,
                "unique_count": data["unique_count"],
                "copy_count": data["copy_count"],
            }
        )

    report = report_payload(source_label, summaries, validator)
    for filename, data in exported:
        write_json(output_dir / filename, data)
    write_json(output_dir / "card_validation.json", report)
    write_markdown_report(output_dir / "card_validation.md", report)

    print(
        f"Export {report['status']}: "
        f"{sum(item['unique_count'] for item in summaries)} definitions, "
        f"{report['error_count']} errors, {report['warning_count']} warnings."
    )
    for item in summaries:
        print(
            f"  {item['pool']}: {item['unique_count']} definitions, "
            f"{item['copy_count']} copies -> {item['filename']}"
        )
    return 1 if report["error_count"] else 0


if __name__ == "__main__":
    raise SystemExit(main())
