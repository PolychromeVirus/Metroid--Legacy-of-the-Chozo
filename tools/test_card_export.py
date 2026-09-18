"""Regression checks for the canonical workbook export."""

from __future__ import annotations

import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path


REPO_ROOT = Path(__file__).resolve().parents[1]
EXPORTER = REPO_ROOT / "tools" / "export_card_data.py"
WORKBOOK = REPO_ROOT / "chozoreference.xlsx"


class CardExportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.temporary = tempfile.TemporaryDirectory()
        cls.output = Path(cls.temporary.name)
        result = subprocess.run(
            [
                sys.executable,
                str(EXPORTER),
                "--workbook",
                str(WORKBOOK),
                "--output",
                str(cls.output),
            ],
            capture_output=True,
            text=True,
            check=False,
        )
        if result.returncode != 0:
            raise AssertionError(
                f"Exporter failed ({result.returncode}).\n"
                f"stdout:\n{result.stdout}\nstderr:\n{result.stderr}"
            )

    @classmethod
    def tearDownClass(cls) -> None:
        cls.temporary.cleanup()

    def load(self, filename: str) -> dict:
        return json.loads((self.output / filename).read_text(encoding="utf-8"))

    def test_expected_pool_totals(self) -> None:
        expected = {
            "cards_loc.json": (50, 131),
            "cards_lop.json": (10, 24),
            "cards_starter.json": (8, 10),
            "metroids.json": (6, 60),
        }
        for filename, totals in expected.items():
            with self.subTest(filename=filename):
                data = self.load(filename)
                self.assertEqual(
                    (data["unique_count"], data["copy_count"]), totals
                )

    def test_ids_are_unique_within_each_pool(self) -> None:
        for filename, key in (
            ("cards_loc.json", "cards"),
            ("cards_lop.json", "cards"),
            ("cards_starter.json", "cards"),
            ("metroids.json", "metroids"),
        ):
            with self.subTest(filename=filename):
                records = self.load(filename)[key]
                ids = [record["definition_id"] for record in records]
                self.assertEqual(len(ids), len(set(ids)))

    def test_relic_type_exports_without_combat_stat(self) -> None:
        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as directory:
            fixture = Path(directory) / "relic.xlsx"
            workbook = load_workbook(WORKBOOK)
            sheet = workbook["Cards"]
            headers = {cell.value: cell.column for cell in sheet[1]}
            sheet.cell(2, headers["type"], "Relic")
            sheet.cell(2, headers["strength"]).value = None
            name = sheet.cell(2, headers["name"]).value
            workbook.save(fixture)
            workbook.close()
            output = Path(directory) / "export"
            result = subprocess.run(
                [sys.executable, str(EXPORTER), "--workbook", str(fixture),
                 "--output", str(output)], capture_output=True, text=True,
            )
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            cards = json.loads((output / "cards_loc.json").read_text())["cards"]
            relic = next(card for card in cards if card["name"] == name)
            self.assertEqual(relic["type"], "relic")
            self.assertIsNone(relic["stat"]["kind"])

    def test_starter_count_header_is_recovered(self) -> None:
        data = self.load("cards_starter.json")
        counts = {card["name"]: card["count"] for card in data["cards"]}
        self.assertEqual(counts["Orders Received"], 2)
        self.assertEqual(counts["Private Military"], 2)
        self.assertEqual(sum(counts.values()), 10)

    def test_fractional_research_values_are_preserved(self) -> None:
        data = self.load("metroids.json")
        values = {
            metroid["name"]: metroid["research_value"]
            for metroid in data["metroids"]
        }
        self.assertEqual(values["Larva"], 0.5)
        self.assertEqual(values["Alpha"], 1.5)

    def test_validation_has_no_errors(self) -> None:
        report = self.load("card_validation.json")
        self.assertEqual(report["status"], "passed")
        self.assertEqual(report["error_count"], 0)
        self.assertEqual(report["warning_count"], 1)
        self.assertEqual(report["issues"][0]["code"], "recovered_header")


if __name__ == "__main__":
    unittest.main()
